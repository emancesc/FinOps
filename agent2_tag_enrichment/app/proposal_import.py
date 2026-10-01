"""
Import di una proposta di tagging già compilata esternamente (xlsx).

Formati accettati (si usa il primo foglio con una colonna ARN che contiene valori):
  - largo: una riga per risorsa, come il foglio "Proposta" dell'export; per ogni tag della
    strategy la colonna "<tag> (proposto)" oppure, se non ce ne sono, la colonna "<tag>";
  - lungo: una riga per proposta, come il foglio "Dettaglio proposte": colonne Arn, Tag e
    "Valore proposto" (o Valore / Value).
Ogni valore è validato contro i valori ammessi della strategy (i non ammessi restano visibili
ma segnalati). Celle vuote, risorse fuori dal job, tag non della strategy e valori uguali a
quello attuale non generano proposte; le proposte già revisionate non vengono sovrascritte.
"""
from __future__ import annotations

import io
from collections import Counter
from typing import Optional

from . import proposal_db as pdb
from .proposal import validate_value

ARN_COLUMNS = {"arn", "resourcearn", "resource arn", "resource_id", "resource id", "resourceid arn"}
TAG_COLUMNS = {"tag", "tag_key", "tag key", "chiave"}
VALUE_COLUMNS = {"valore proposto", "valore", "value", "tag_value", "proposed value", "proposto"}
HEADER_SEARCH_ROWS = 5


def _norm(value) -> str:
    return str(value or "").strip().lower()


def _text(value) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def read_entries(content: bytes, tag_keys: list[str]) -> tuple[str, list[tuple[int, str, str, str]]]:
    """Ritorna (foglio usato, [(riga excel, arn, tag, valore)])."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    proposed_cols = {f"{k} (proposto)".lower(): k for k in tag_keys}
    plain_cols = {k.lower(): k for k in tag_keys}
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header, offset = None, 0
        for offset in range(1, HEADER_SEARCH_ROWS + 1):
            row = next(rows, None)
            if row is None:
                break
            if any(_norm(c) in ARN_COLUMNS for c in row):
                header = [_norm(c) for c in row]
                break
        if not header:
            continue
        arn_i = next(i for i, c in enumerate(header) if c in ARN_COLUMNS)
        tag_i = next((i for i, c in enumerate(header) if c in TAG_COLUMNS), None)
        value_i = next((i for i, c in enumerate(header) if c in VALUE_COLUMNS), None)
        wide = {i: proposed_cols[c] for i, c in enumerate(header) if c in proposed_cols} \
            or {i: plain_cols[c] for i, c in enumerate(header) if c in plain_cols}
        entries = []
        for n, row in enumerate(rows, start=offset + 1):
            arn = _text(row[arn_i]) if arn_i < len(row) else None
            if not arn:
                continue
            if tag_i is not None and value_i is not None:
                tag = _text(row[tag_i]) if tag_i < len(row) else None
                value = _text(row[value_i]) if value_i < len(row) else None
                if tag and value:
                    entries.append((n, arn, tag, value))
            else:
                for i, key in wide.items():
                    value = _text(row[i]) if i < len(row) else None
                    if value:
                        entries.append((n, arn, key, value))
        if entries:
            return ws.title, entries
    raise ValueError("Nessun foglio con una colonna ARN (Arn / ResourceArn / resource_id) e valori di tag: "
                     "usare il formato dell'export (foglio Proposta o Dettaglio proposte)")


def import_xlsx(job_id: str, file_name: str, content: bytes, approve: bool = False,
                reviewed_by: Optional[str] = "operator") -> dict:
    job = pdb.get_job(job_id)
    if not job:
        raise LookupError(f"Job {job_id} non trovato")
    strategy = pdb.get_strategy(None)
    if not strategy:
        raise ValueError("Nessuna Tagging Strategy attiva (pagina Regole)")
    tags_by_key = {t["tag_key"]: t for t in strategy["tags"]}
    sheet, entries = read_entries(content, list(tags_by_key))
    current_tags = pdb.load_resource_tags(job_id)

    stats = Counter()
    unknown_resources, unknown_tags, proposals = set(), set(), {}
    for row, arn, key, value in entries:
        if arn not in current_tags:
            unknown_resources.add(arn)
            continue
        if key not in tags_by_key:
            unknown_tags.add(key)
            continue
        current = (current_tags[arn] or {}).get(key)
        if current == value:
            stats["unchanged"] += 1
            continue
        ok, problem = validate_value(tags_by_key[key], value)
        stats["invalid"] += not ok
        reasoning = f"Importato da {file_name}"
        proposals[(arn, key)] = {  # a parità di risorsa e tag vale l'ultima riga del file
            "resource_id": arn, "tag_key": key, "tag_value": value, "current_value": current,
            "confidence": 1.0 if ok else 0.2, "source_type": "manual_override",
            "source_ref": f"{file_name} / {sheet} riga {row}",
            "reasoning": reasoning if ok else f"[VALORE NON AMMESSO: {problem}] {reasoning}"}

    options = {"source": "xlsx", "file": file_name, "sheet": sheet, "approve": approve}
    run = pdb.create_run(job_id, strategy["strategy_id"], options, len(current_tags))
    saved = pdb.save_run_proposals(job_id, run["run_id"], strategy["strategy_id"], list(proposals.values()),
                                   review_status="approved" if approve else "pending", reviewed_by=reviewed_by)
    resources = len({arn for arn, _ in proposals})
    summary = {
        "run_id": run["run_id"], "sheet": sheet, "entries": len(entries), "proposals": len(proposals),
        "saved": saved, "not_overwritten_reviewed": len(proposals) - saved, "resources": resources,
        "unchanged": stats["unchanged"], "invalid_values": stats["invalid"],
        "unknown_resources": len(unknown_resources), "unknown_resources_sample": sorted(unknown_resources)[:10],
        "unknown_tags": sorted(unknown_tags), "review_status": "approved" if approve else "pending",
    }
    message = (f"Import {file_name} ({sheet}): {saved} proposte su {resources} risorse"
               + (f", {stats['invalid']} valori non ammessi" if stats["invalid"] else "")
               + (f", {len(unknown_resources)} ARN fuori dal job" if unknown_resources else "")
               + (f", {len(proposals) - saved} già revisionate non sovrascritte" if len(proposals) - saved else ""))
    pdb.update_run(run["run_id"], status="done", started=True, finished=True, progress_pct=100.0,
                   resources_done=resources, proposals_saved=saved, message=message)
    return summary
