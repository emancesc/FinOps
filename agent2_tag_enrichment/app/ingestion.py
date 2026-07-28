"""
Ingestion documentale: parsing PDF/DOCX/XLSX, chunking, embedding, persistenza.
"""
from __future__ import annotations
import logging
import os
from typing import Callable

from .embedding import embed as default_embed
from .db import save_chunks, upsert_document, mark_document_parsed

logger = logging.getLogger(__name__)

CHUNK_SIZE = 500   # caratteri
CHUNK_OVERLAP = 100


def ingest(
    job_id: str,
    document_id: str,
    storage_path: str,
    doc_type: str,
    embed_fn: Callable | None = None,
) -> int:
    """
    Parsa il documento, fa chunking, genera embedding e persiste in document_chunks.
    Ritorna il numero di chunk salvati.
    """
    _embed = embed_fn or default_embed
    file_name = os.path.basename(storage_path)

    upsert_document(job_id, document_id, doc_type, file_name, storage_path)

    ext = os.path.splitext(storage_path)[1].lower()
    if ext == ".pdf":
        text = _parse_pdf(storage_path)
    elif ext in (".docx", ".doc"):
        text = _parse_docx(storage_path)
    elif ext in (".xlsx", ".xls"):
        text = _parse_xlsx(storage_path)
    else:
        text = _parse_text(storage_path)

    raw_chunks = _chunk(text)
    if not raw_chunks:
        logger.warning("Documento %s non ha prodotto chunk", storage_path)
        return 0

    chunk_tuples = []
    for idx, content in enumerate(raw_chunks):
        embedding = _embed(content)
        meta = {"chunk_index": idx, "doc_type": doc_type, "source": file_name}
        chunk_tuples.append((idx, content, embedding, meta))

    save_chunks(document_id, chunk_tuples)
    mark_document_parsed(document_id)
    logger.info("Documento %s: %d chunk salvati", document_id, len(chunk_tuples))
    return len(chunk_tuples)


def _parse_pdf(path: str) -> str:
    from pypdf import PdfReader
    reader = PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _parse_docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _parse_xlsx(path: str) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            line = "\t".join(str(c) if c is not None else "" for c in row)
            if line.strip():
                lines.append(line)
    return "\n".join(lines)


def _parse_text(path: str) -> str:
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read()


def _chunk(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Chunking semplice a finestra scorrevole su caratteri."""
    text = text.strip()
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start += size - overlap
    return chunks
