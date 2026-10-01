"""
Export xlsx della proposta di tagging, con colonne simili all'export di Resource Explorer
(scripts/export_resource_explorer_xlsx.py):
  - "Proposta": una riga per risorsa del job, con per ogni tag della strategy il valore
    attuale e quello proposto affiancati, e la conformità ai tag obbligatori prima e dopo;
  - "Dettaglio proposte": una riga per proposta (valori, confidenza, fonte, motivazione, revisione);
  - "Summary": job, strategy, run e copertura per tag.
Le proposte rifiutate compaiono solo nel dettaglio, non come valore proposto.
"""
from __future__ import annotations

import io
import json
from collections import Counter
from datetime import datetime, timezone
from typing import Optional

MAX_CELL = 32000  # limite Excel: 32767 caratteri per cella
FIXED_COLUMNS = ["Region", "Service", "ResourceType", "Arn", "ResourceId", "Name", "OwningAccountId",
                 "cineca_mandatory_compliant", "cineca_mandatory_compliant_proposto", "cineca_mandatory_missing_proposto",
                 "proposte", "revisione"]
DETAIL_COLUMNS = ["Arn", "ResourceType", "Region", "Tag", "Valore attuale", "Valore proposto", "Confidenza", "Fonte",
                  "Riferimento", "Motivazione", "Stato revisione", "Revisionato da", "Aggiornato"]


def _cell(value):
    if value is None:
        return None
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    value = str(value)
    return value if len(value) <= MAX_CELL else value[:MAX_CELL] + "...[truncated]"


def _service(arn: str) -> Optional[str]:
    parts = arn.split(":")
    return parts[2] if len(parts) > 2 and parts[0] == "arn" else None


def _resource_id(arn: str) -> str:
    tail = arn.split(":", 5)[-1] if arn.count(":") >= 5 else arn
    return tail.split("/", 1)[-1] if "/" in tail else tail


def _col_letter(n: int) -> str:
    s = ""
    while n:
        n, rem = divmod(n - 1, 26)
        s = chr(65 + rem) + s
    return s


def build_xlsx(job: dict, strategy: dict, run_id: Optional[str], resources: list[dict], proposals: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill

    tag_keys = [t["tag_key"] for t in strategy["tags"]]
    mandatory = [t["tag_key"] for t in strategy["tags"] if t.get("mandatory")]
    by_resource: dict[str, dict[str, dict]] = {}
    for p in proposals:
        by_resource.setdefault(p["resource_id"], {})[p["tag_key"]] = p

    wb = Workbook(write_only=True)
    bold = Font(bold=True, color="FFFFFF")
    fills = {"fixed": PatternFill("solid", fgColor="1F3864"), "current": PatternFill("solid", fgColor="7B2C8C"),
             "proposed": PatternFill("solid", fgColor="2E7D32")}

    def header_row(ws, names, kinds=None):
        cells = []
        for i, name in enumerate(names):
            c = WriteOnlyCell(ws, value=name)
            c.font, c.fill = bold, fills[(kinds or {}).get(i, "fixed")]
            cells.append(c)
        ws.append(cells)

    # --- Proposta: una riga per risorsa ------------------------------------
    header = list(FIXED_COLUMNS)
    kinds = {}
    for key in tag_keys:
        kinds[len(header)] = "current"
        header.append(key)
        kinds[len(header)] = "proposed"
        header.append(f"{key} (proposto)")
    header.append("Tags (JSON)")
    ws = wb.create_sheet("Proposta")
    ws.freeze_panes = "G2"
    ws.auto_filter.ref = f"A1:{_col_letter(len(header))}{len(resources) + 1}"
    header_row(ws, header, kinds)

    coverage_now, coverage_after, statuses = Counter(), Counter(), Counter()
    compliant_now = compliant_after = 0
    for res in resources:
        arn = res["resource_id"]
        tags = res.get("current_tags") or {}
        props = by_resource.get(arn, {})
        proposed = {k: p["tag_value"] for k, p in props.items() if p["review_status"] != "rejected" and p.get("tag_value")}
        after = {**{k: v for k, v in tags.items() if v}, **proposed}
        missing_now = [k for k in mandatory if not tags.get(k)]
        missing_after = [k for k in mandatory if not after.get(k)]
        compliant_now += not missing_now
        compliant_after += not missing_after
        coverage_now.update(k for k in tag_keys if tags.get(k))
        coverage_after.update(k for k in tag_keys if after.get(k))
        review = Counter(p["review_status"] for p in props.values())
        statuses.update(review)
        row = [res.get("region"), _service(arn), res.get("resource_type"), arn, _resource_id(arn), tags.get("Name"),
               job["account_id"], "SI" if not missing_now else "NO", "SI" if not missing_after else "NO",
               ", ".join(missing_after) or None, len(props) or None,
               ", ".join(f"{s} {n}" for s, n in sorted(review.items())) or None]
        for key in tag_keys:
            row += [_cell(tags.get(key)), _cell(proposed.get(key))]
        row.append(_cell(tags) if tags else None)
        ws.append(row)

    # --- Dettaglio proposte -----------------------------------------------
    ws = wb.create_sheet("Dettaglio proposte")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{_col_letter(len(DETAIL_COLUMNS))}{len(proposals) + 1}"
    header_row(ws, DETAIL_COLUMNS)
    for p in proposals:
        ws.append([p["resource_id"], p.get("resource_type"), p.get("region"), p["tag_key"], _cell(p.get("current_value")),
                   _cell(p.get("tag_value")), float(p["confidence"]) if p.get("confidence") is not None else None,
                   p.get("source_type"), _cell(p.get("source_ref")), _cell(p.get("reasoning")), p.get("review_status"),
                   p.get("reviewed_by"), str(p.get("updated_at") or "")[:19] or None])

    # --- Summary ----------------------------------------------------------
    total = len(resources)
    pct = lambda n: round(100.0 * n / total, 1) if total else 0.0  # noqa: E731
    ws = wb.create_sheet("Summary")
    for row in (["Tenant", job.get("tenant_id")], ["Account", job["account_id"]], ["Job", str(job["job_id"])],
                ["Tagging Strategy", f"{strategy['name']} rev. {strategy['revision']}"],
                ["Esecuzione", run_id or "tutte"], ["Generato", datetime.now(timezone.utc).isoformat(timespec="seconds")],
                ["Risorse", total], ["Proposte", len(proposals)],
                ["Conformi ai tag obbligatori (attuale)", compliant_now], ["% conformi (attuale)", pct(compliant_now)],
                ["Conformi ai tag obbligatori (con proposta)", compliant_after],
                ["% conformi (con proposta)", pct(compliant_after)], []):
        ws.append(row)
    header_row(ws, ["Tag", "obbligatorio", "risorse con tag (attuale)", "% attuale", "risorse con tag (con proposta)",
                    "% con proposta"])
    for key in tag_keys:
        ws.append([key, "SI" if key in mandatory else "NO", coverage_now[key], pct(coverage_now[key]),
                   coverage_after[key], pct(coverage_after[key])])
    ws.append([])
    header_row(ws, ["Stato revisione", "proposte"])
    for status, n in sorted(statuses.items()):
        ws.append([status, n])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
