"""
Estrazione via LLM delle regole di una Tagging Strategy da un documento (PDF/DOCX/testo).

Il documento viene letto pagina per pagina e diviso in blocchi; per ogni blocco
l'LLM restituisce, in JSON validato, i metadati (nome, revisione, data di
rilascio, changelog), la definizione dei tag (categoria, obbligatorietà,
multi-valore, valori ammessi) e le regole applicative (vincoli, regole per tipo
di risorsa, pattern di naming, risorse condivise, esempi reali).

Ogni blocco viene salvato nel DB appena elaborato (strategy_tags / strategy_rules
e avanzamento su tagging_strategies): un'estrazione interrotta riprende dal
blocco successivo all'ultimo salvato.
"""
from __future__ import annotations

import json
import logging
import os
from typing import Optional

from pydantic import BaseModel, Field

from . import strategy_db as sdb

logger = logging.getLogger(__name__)

CHUNK_CHARS = int(os.environ.get("STRATEGY_CHUNK_CHARS", "9000"))
MAX_OUTPUT_TOKENS = int(os.environ.get("STRATEGY_MAX_OUTPUT_TOKENS", "32000"))


# ---------------------------------------------------------------------------
# Lettura documento
# ---------------------------------------------------------------------------

def read_pages(path: str) -> list[tuple[int, str]]:
    """Ritorna [(numero_pagina, testo)]. Per DOCX/testo una 'pagina' ogni ~3000 caratteri."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        return [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
    if ext in (".docx", ".doc"):
        from docx import Document
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:  # le tabelle dei valori ammessi sono il contenuto principale
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
    else:
        with open(path, encoding="utf-8", errors="replace") as f:
            text = f.read()
    size = 3000
    return [(i // size + 1, text[i:i + size]) for i in range(0, len(text), size)] or [(1, "")]


def make_chunks(pages: list[tuple[int, str]], max_chars: int = CHUNK_CHARS) -> list[dict]:
    """Raggruppa pagine consecutive in blocchi di al massimo ~max_chars caratteri."""
    chunks, current, first = [], [], None
    for number, text in pages:
        text = text.strip()
        if not text:
            continue
        if current and sum(len(t) for t in current) + len(text) > max_chars:
            chunks.append({"pages": f"{first}-{prev}" if first != prev else str(first), "text": "\n".join(current)})
            current, first = [], None
        if first is None:
            first = number
        current.append(f"[pagina {number}]\n{text}")
        prev = number
    if current:
        chunks.append({"pages": f"{first}-{prev}" if first != prev else str(first), "text": "\n".join(current)})
    return chunks


# ---------------------------------------------------------------------------
# Schema della risposta LLM
# ---------------------------------------------------------------------------

class ChangelogEntry(BaseModel):
    version: str
    date: Optional[str] = None
    authors: Optional[str] = None
    changes: list[str] = []


class StrategyMeta(BaseModel):
    name: Optional[str] = None
    revision: Optional[str] = None
    release_date: Optional[str] = Field(default=None, description="YYYY-MM-DD")
    summary: Optional[str] = None
    changelog: list[ChangelogEntry] = []


class AllowedValue(BaseModel):
    value: str
    description: Optional[str] = None
    business_unit: Optional[str] = None


class TagDef(BaseModel):
    tag_key: str
    category: str = "other"          # cost_allocation | operational | other
    description: Optional[str] = None
    mandatory: bool = False
    billing: bool = False
    multi_value: bool = False
    separator: Optional[str] = None
    allowed_values: list[AllowedValue] = []
    notes: list[str] = []
    source_ref: Optional[str] = None


class RuleDef(BaseModel):
    rule_type: str = "guideline"     # guideline | resource_type | naming_pattern | shared_resource | value_constraint | example | other
    tag_keys: list[str] = []
    title: str
    description: str
    condition: dict = {}
    resolution: dict = {}
    source_ref: Optional[str] = None


class ChunkExtraction(BaseModel):
    meta: Optional[StrategyMeta] = None
    tags: list[TagDef] = []
    rules: list[RuleDef] = []
    last_tag_in_progress: Optional[str] = None


SYSTEM_PROMPT = """Sei un analista FinOps che converte un documento di Tagging Strategy AWS in regole strutturate.
Ricevi UN BLOCCO del documento (le pagine sono marcate con [pagina N]). Estrai SOLO ciò che è scritto
nel blocco, senza inventare nulla e senza riassumere le liste.

Estrai:
1. meta (solo se il blocco li contiene): nome del documento, numero di revisione (es. "1.6"), data di
   rilascio della revisione corrente in formato YYYY-MM-DD, breve sintesi, changelog (una voce per versione).
2. tags: per ogni tag definito o i cui valori ammessi compaiono nel blocco:
   - tag_key esattamente come scritto (con il prefisso, es. "cineca:Customer");
   - category: "cost_allocation" per i tag di allocazione costi/billing, "operational" per quelli operativi;
   - mandatory, billing (true se vuoto o N/A non sono ammessi), multi_value, separator (es. "+");
   - allowed_values: TUTTE le righe della tabella dei valori ammessi presenti nel blocco, con "value"
     copiato ESATTAMENTE come scritto (maiuscole, spazi, punti, apostrofi), descrizione e business unit se presente;
   - notes: vincoli e indicazioni specifiche del tag (una frase ciascuno); source_ref: sezione e pagina.
   Se il blocco inizia continuando una tabella di un tag del blocco precedente (indicato dall'utente come
   "tag in corso"), attribuisci quelle righe a quel tag.
3. rules: regole applicabili nella proposta di tagging:
   - "value_constraint": vincoli su valori (es. separatore, valori combinati ammessi, vuoto non ammesso);
   - "resource_type": regole per tipi di risorsa (condition.resource_types con i nomi citati,
     resolution con il valore da assegnare, anche stringa vuota se il tag va lasciato vuoto);
   - "naming_pattern": pattern di nomi risorsa (condition.name_patterns) con i valori da assegnare;
   - "shared_resource": regole per risorse condivise (condition.max_consumers ecc.);
   - "example": ogni scenario reale descritto (condition con tipo/nome della risorsa, resolution con
     la mappa completa tag -> valore dell'esempio);
   - "guideline": principi generali. Una regola per concetto, con source_ref (sezione, pagina).
4. last_tag_in_progress: il tag_key la cui tabella dei valori ammessi continua oltre la fine del blocco, altrimenti null.

Rispondi SOLO con un oggetto JSON con le chiavi meta, tags, rules, last_tag_in_progress."""


# ---------------------------------------------------------------------------
# Merge dei risultati nel DB
# ---------------------------------------------------------------------------

def _merge_tag(existing: Optional[dict], new: TagDef) -> dict:
    if not existing:
        return new.model_dump()
    merged = dict(existing)
    for flag in ("mandatory", "billing", "multi_value"):
        merged[flag] = bool(existing.get(flag)) or getattr(new, flag)
    if new.category != "other":
        merged["category"] = new.category
    for field in ("description", "separator", "source_ref"):
        if getattr(new, field) and len(str(getattr(new, field))) > len(str(existing.get(field) or "")):
            merged[field] = getattr(new, field)
    values = {v["value"]: v for v in existing.get("allowed_values") or []}
    for v in new.allowed_values:
        values.setdefault(v.value, v.model_dump())
    merged["allowed_values"] = list(values.values())
    merged["notes"] = list(dict.fromkeys((existing.get("notes") or []) + new.notes))
    return merged


async def extract_strategy(strategy_id: str, llm_client=None, fresh: bool = False) -> dict:
    """Estrae (o riprende l'estrazione di) una strategy. Ritorna la strategy aggiornata."""
    from llm_gateway.base import LLMMessage

    if llm_client is None:
        from llm_gateway.factory import get_llm_client
        llm_client = get_llm_client()

    strategy = sdb.get_strategy(strategy_id)
    if not strategy:
        raise ValueError(f"Strategy {strategy_id} non trovata")
    if fresh:
        sdb.reset_extraction(strategy_id)
        strategy = sdb.get_strategy(strategy_id)

    chunks = make_chunks(read_pages(strategy["storage_path"]))
    start = min(int(strategy.get("chunks_done") or 0), len(chunks))
    sdb.update_progress(strategy_id, status="extracting", chunks_total=len(chunks), chunks_done=start,
                        llm_model=getattr(llm_client, "_model", None), error=None)
    logger.info("Strategy %s: %d blocchi, ripresa da %d", strategy_id, len(chunks), start)

    tag_in_progress = strategy.get("_tag_in_progress")
    for index in range(start, len(chunks)):
        chunk = chunks[index]
        user = (
            f"Blocco {index + 1} di {len(chunks)} (pagine {chunk['pages']}).\n"
            f"Tag in corso dal blocco precedente: {tag_in_progress or 'nessuno'}\n\n{chunk['text']}"
        )
        try:
            resp = await llm_client.complete(
                system=SYSTEM_PROMPT,
                messages=[LLMMessage(role="user", content=user)],
                response_format=ChunkExtraction,
                max_tokens=MAX_OUTPUT_TOKENS,
            )
            from llm_gateway.claude_client import _strip_fence
            result = ChunkExtraction.model_validate(json.loads(_strip_fence(resp.content)))
        except Exception as exc:  # noqa: BLE001
            logger.exception("Estrazione strategy %s, blocco %d fallita", strategy_id, index + 1)
            sdb.update_progress(strategy_id, status="error", error=f"blocco {index + 1}: {exc}")
            return sdb.get_strategy(strategy_id)

        for tag in result.tags:
            existing = sdb.get_tag(strategy_id, tag.tag_key)
            sdb.upsert_tag(strategy_id, _merge_tag(existing, tag))
        for rule in result.rules:
            sdb.insert_rule(strategy_id, rule.model_dump())
        if result.meta:
            sdb.merge_meta(strategy_id, result.meta.model_dump())
        tag_in_progress = result.last_tag_in_progress
        sdb.update_progress(strategy_id, chunks_done=index + 1,
                            progress_pct=round(100.0 * (index + 1) / len(chunks), 1),
                            tag_in_progress=tag_in_progress)
        logger.info("Strategy %s: blocco %d/%d salvato (%d tag, %d regole)",
                    strategy_id, index + 1, len(chunks), len(result.tags), len(result.rules))

    sdb.update_progress(strategy_id, status="extracted", progress_pct=100.0, finished=True)
    return sdb.get_strategy(strategy_id)
