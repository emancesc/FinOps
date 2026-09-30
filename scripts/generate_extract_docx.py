"""Genera il docx "Estrazione file JSON" per un account AWS qualsiasi.

Uso:
  python scripts/generate_extract_docx.py --account <id> --profile <profilo>
      [--label "Nome account"] [--regions all] [--output file.docx]

Gli script inclusi nel documento estraggono tutte le regioni abilitate
dell'account (o una lista di regioni) e salvano gli oggetti AWS completi:
nessuna proprieta' viene scartata. Lo script PowerShell incluso e' il
contenuto integrale di scripts/extract_linked_resources.ps1.
"""
import argparse
import os

from docx import Document

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _powershell_script(account: str, profile: str, regions: str) -> list[str]:
    """Comando di lancio + contenuto integrale di extract_linked_resources.ps1."""
    with open(os.path.join(REPO, "scripts", "extract_linked_resources.ps1"), encoding="utf-8") as f:
        body = f.read().splitlines()
    return [
        "# Lancio (produce tutti i file di evidenza per ogni regione con risorse):",
        f".\scripts\extract_linked_resources.ps1 -AwsProfile {profile} -Account {account} -Regions {regions}",
        "",
        "# Contenuto di scripts\extract_linked_resources.ps1:",
        *body,
    ]


def _python_section(account: str, profile: str, regions: str) -> list[str]:
    return [
        "# Estrazione completa (SSM, CloudFormation, EIP, ACM, Config rules, ENI, volumi,",
        "# snapshot, istanze + volumes_report.json) per tutte le regioni con risorse:",
        f"python {REPO}\\scripts\\extract_linked_resources.py --profile {profile} --account {account} --regions {regions}",
        "",
        "# Verifica output per regione",
        "import os",
        f"root = r'{REPO}\\extracted\\{account}'",
        "for region in sorted(os.listdir(root)):",
        "    region_dir = os.path.join(root, region)",
        "    if not os.path.isdir(region_dir):",
        "        continue",
        "    for name in sorted(os.listdir(region_dir)):",
        "        if name.endswith('.json'):",
        "            print(f'{region:15s} {name:40s} size={os.path.getsize(os.path.join(region_dir, name))}')",
    ]


def build_extract_docx(account: str, profile: str, title: str, output_path: str, regions: str = "all") -> str:
    sections = [
        ("1. Autenticazione AWS SSO", [
            f"aws sso login --profile {profile}",
            f"aws sts get-caller-identity --profile {profile} --output json",
        ]),
        ("2. Script PowerShell per estrazione dei file JSON (tutte le evidenze, multi-regione)",
         _powershell_script(account, profile, regions)),
        ("3. Script Python: estrazione completa e verifica output",
         _python_section(account, profile, regions)),
    ]

    doc = Document()
    doc.add_heading(title, 0)
    doc.add_paragraph(f"Account: {account}")
    doc.add_paragraph(f"Regioni: {regions} (output in extracted\\{account}\\<regione>, solo regioni con risorse)")
    doc.add_paragraph(f"Profile: {profile}")

    for heading, lines in sections:
        doc.add_heading(heading, level=1)
        for line in lines:
            p = doc.add_paragraph()
            if not line.strip():
                p.add_run("\n")
                continue
            p.add_run(line)

    doc.save(output_path)
    return output_path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Genera il docx con gli script di estrazione JSON per un account")
    parser.add_argument("--account", required=True, help="account id AWS")
    parser.add_argument("--profile", required=True, help="profilo AWS CLI (SSO)")
    parser.add_argument("--label", default=None, help="nome leggibile dell'account nel titolo (default: account id)")
    parser.add_argument("--regions", default="all", help='"all" (default), una regione o lista "a,b"')
    parser.add_argument("--output", default=None, help="default: extract_scripts_<account>.docx nella root del repo")
    args = parser.parse_args()

    path = build_extract_docx(
        account=args.account,
        profile=args.profile,
        title=f"Estrazione file JSON account {args.label or args.account}",
        output_path=args.output or os.path.join(REPO, f"extract_scripts_{args.account}.docx"),
        regions=args.regions,
    )
    print(path)
