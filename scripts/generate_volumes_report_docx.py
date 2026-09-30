"""Genera il docx "Volumi EBS e collegamenti ad altre risorse" per un account qualsiasi.

Uso:
  python scripts/generate_volumes_report_docx.py --account <id>
      [--label "Nome account"] [--profile <profilo>] [--extracted-root extracted]
      [--notes-file note.txt] [--output file.docx]

Legge extracted/<account>/<regione>/volumes_report.json di tutte le regioni
estratte da extract_linked_resources.py. Le note (analisi manuale) sono
opzionali: un paragrafo per blocco di testo separato da una riga vuota.
"""
import argparse
import json
import os

from docx import Document

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def load_volumes(account_dir: str) -> tuple[list[dict], list[str]]:
    volumes: list[dict] = []
    regions: list[str] = []
    for region in sorted(os.listdir(account_dir)):
        report = os.path.join(account_dir, region, "volumes_report.json")
        if not os.path.isfile(report):
            continue
        regions.append(region)
        with open(report, encoding="utf-8") as f:
            for v in json.load(f):
                v["Region"] = region
                volumes.append(v)
    volumes.sort(key=lambda v: (v["Region"], v["Attachments"][0]["InstanceName"] or "" if v["Attachments"] else "zzz_non_attaccato", v["VolumeId"]))
    return volumes, regions


def build_volumes_report_docx(
    account: str,
    output_path: str,
    label: str | None = None,
    profile: str | None = None,
    extracted_root: str = os.path.join(REPO, "extracted"),
    notes: list[str] | None = None,
) -> str:
    volumes, regions = load_volumes(os.path.join(extracted_root, account))
    if not volumes:
        raise SystemExit(f"Nessun volumes_report.json in {os.path.join(extracted_root, account)}")

    doc = Document()
    doc.add_heading(f"Volumi EBS e collegamenti ad altre risorse - Account {label or account}", 0)
    doc.add_paragraph(f"Account: {account}")
    doc.add_paragraph(f"Regioni: {', '.join(regions)}")
    if profile:
        doc.add_paragraph(f"Profile: {profile}")

    total_size = sum(v["Size"] for v in volumes)
    attached = [v for v in volumes if v["Attachments"]]
    unattached = [v for v in volumes if not v["Attachments"]]
    with_snaps = [v for v in volumes if v["Snapshots"]]

    doc.add_heading("1. Riepilogo", level=1)
    doc.add_paragraph(f"Totale volumi: {len(volumes)}")
    doc.add_paragraph(f"Volumi collegati a un'istanza EC2: {len(attached)}")
    doc.add_paragraph(f"Volumi non collegati: {len(unattached)}")
    doc.add_paragraph(f"Volumi con almeno uno snapshot: {len(with_snaps)}")
    doc.add_paragraph(f"Dimensione totale: {total_size} GiB")
    for region in regions:
        doc.add_paragraph(f"  {region}: {sum(1 for v in volumes if v['Region'] == region)} volumi")

    doc.add_heading("2. Dettaglio volumi", level=1)
    table = doc.add_table(rows=1, cols=9)
    table.style = "Light Grid Accent 1"
    headers = ["Region", "VolumeId", "Size (GiB)", "Type", "Encrypted", "Istanza (Name)", "InstanceId", "Device", "Snapshot"]
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text

    for v in volumes:
        att = v["Attachments"][0] if v["Attachments"] else None
        row = table.add_row().cells
        row[0].text = v["Region"]
        row[1].text = v["VolumeId"]
        row[2].text = str(v["Size"])
        row[3].text = v["VolumeType"]
        row[4].text = "Si" if v["Encrypted"] else "No"
        row[5].text = (att["InstanceName"] or "-") if att else "NON ATTACCATO"
        row[6].text = att["InstanceId"] if att else "-"
        row[7].text = att["Device"] if att else "-"
        row[8].text = v["Snapshots"][0]["SnapshotId"] if v["Snapshots"] else "-"

    if notes:
        doc.add_heading("3. Note", level=1)
        for note in notes:
            doc.add_paragraph(note)

    doc.save(output_path)
    return output_path


def read_notes(path: str) -> list[str]:
    with open(path, encoding="utf-8") as f:
        blocks = f.read().split("\n\n")
    return [" ".join(b.split()) for b in blocks if b.strip()]


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Genera il docx dei volumi EBS di un account")
    parser.add_argument("--account", required=True, help="account id AWS (cartella in extracted/)")
    parser.add_argument("--label", default=None, help="nome leggibile dell'account nel titolo")
    parser.add_argument("--profile", default=None, help="profilo AWS usato per l'estrazione (solo informativo)")
    parser.add_argument("--extracted-root", default=os.path.join(REPO, "extracted"))
    parser.add_argument("--notes-file", default=None, help="default: <extracted-root>/<account>/volumes_report_notes.txt se esiste")
    parser.add_argument("--output", default=None, help="default: volumes_report_<account>.docx nella root del repo")
    args = parser.parse_args()

    notes_file = args.notes_file or os.path.join(args.extracted_root, args.account, "volumes_report_notes.txt")
    path = build_volumes_report_docx(
        account=args.account,
        output_path=args.output or os.path.join(REPO, f"volumes_report_{args.account}.docx"),
        label=args.label,
        profile=args.profile,
        extracted_root=args.extracted_root,
        notes=read_notes(notes_file) if os.path.isfile(notes_file) else None,
    )
    print(path)
