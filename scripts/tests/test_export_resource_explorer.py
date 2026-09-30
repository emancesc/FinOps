"""
Test offline di scripts/export_resource_explorer_xlsx.py (AWS finto):
- un unico xlsx per tenant con la regione di estrazione di ogni risorsa
- tutti i campi, una colonna per ogni tag cineca:* (ordine strategy) e per ogni altro tag
- paginazione persistita pagina per pagina e ripresa dopo un'interruzione
- regione negata da SCP marcata "denied" e non ritentata
"""
import importlib.util
import json
import os
import sys

import pytest
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPT = os.path.join(os.path.dirname(HERE), "export_resource_explorer_xlsx.py")
ACCOUNT = "123456789012"

sys.path.insert(0, HERE)
from fake_aws_data import fake_aws  # noqa: E402


@pytest.fixture
def mod(monkeypatch):
    spec = importlib.util.spec_from_file_location("export_resource_explorer_xlsx", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    monkeypatch.setattr(module.elr, "aws_json", fake_aws)
    return module


def _sheet(path, name):
    ws = load_workbook(path, read_only=True)[name]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    return rows[0], [dict(zip(rows[0], r)) for r in rows[1:]]


def _argv(tmp_path, *extra):
    return ["--profile", "p", "--tenant", "demo", "--out-root", str(tmp_path), *extra]


def test_single_xlsx_with_region_and_all_tags(mod, tmp_path):
    xlsx = mod.main(_argv(tmp_path))
    assert xlsx == str(tmp_path / ACCOUNT / "resource_explorer_demo.xlsx")

    header, rows = _sheet(xlsx, "Resources")
    assert header[:6] == ["ExtractionRegion", "Region", "Service", "ResourceType", "CfnResourceType", "Arn"]
    # colonne cineca: tutte quelle della strategy nell'ordine previsto, poi quelle extra trovate
    cineca = [h for h in header if h.startswith("cineca:") and not h.startswith("cineca_")]
    assert cineca[:4] == ["cineca:BusinessUnit", "cineca:Customer", "cineca:Product", "cineca:Environment"]
    assert "cineca:Custom" in cineca
    assert {"Name", "Owner"} <= set(header) and header[-2:] == ["Tags (JSON)", "OtherProperties (JSON)"]

    assert {r["Region"] for r in rows} == {"eu-south-1", "us-east-1"}
    assert all(r["ExtractionRegion"] == r["Region"] for r in rows)
    assert len(rows) == 5
    by_id = {r["ResourceId"]: r for r in rows}
    inst = by_id["i-1"]
    assert inst["Region"] == "eu-south-1" and inst["Name"] == "web-01" and inst["Owner"] == "team-a"
    assert inst["cineca:Product"] == "ESSE3" and inst["cineca_mandatory_compliant"] == "SI"
    vol = by_id["vol-1"]
    assert vol["cineca_mandatory_compliant"] == "NO"
    assert "cineca:BusinessUnit" in vol["cineca_mandatory_missing"]
    assert json.loads(inst["Tags (JSON)"])["cineca:Role"] == "Web"
    assert by_id["c-1"]["Region"] == "us-east-1"

    _, regions = _sheet(xlsx, "Regions")
    status = {r["Region"]: r["status"] for r in regions}
    assert status == {"eu-south-1": "done", "eu-west-1": "denied", "us-east-1": "done"}

    summary = [list(r) for r in load_workbook(xlsx, read_only=True)["Summary"].iter_rows(values_only=True)]
    assert ["Total resources", 5] in summary
    assert ["cineca mandatory compliant", 1] in summary


def test_resume_from_saved_page(mod, tmp_path, monkeypatch):
    """Interruzione dopo la prima pagina di eu-south-1: la ripresa chiede solo la seconda."""
    class Crash(BaseException):
        pass

    def crashing(region, *args):
        if region == "eu-south-1" and "--next-token" in args:
            raise Crash()
        return fake_aws(region, *args)

    monkeypatch.setattr(mod.elr, "aws_json", crashing)
    monkeypatch.setattr(mod, "MAX_WORKERS", 1)
    with pytest.raises(Crash):
        mod.main(_argv(tmp_path, "--regions", "eu-south-1"))
    data = tmp_path / ACCOUNT / "resource_explorer"
    state = json.loads((data / "eu-south-1.state.json").read_text(encoding="utf-8"))
    assert state["pages"] == 1 and state["next_token"] == "page-1"
    assert len((data / "eu-south-1.jsonl").read_text(encoding="utf-8").splitlines()) == 2

    calls = []

    def recording(region, *args):
        calls.append((region, *args))
        return fake_aws(region, *args)

    monkeypatch.setattr(mod.elr, "aws_json", recording)
    xlsx = mod.main(_argv(tmp_path, "--regions", "eu-south-1"))
    pages = [c for c in calls if c[1:3] == ("resource-explorer-2", "list-resources")]
    assert len(pages) == 1 and "page-1" in pages[0]
    _, rows = _sheet(xlsx, "Resources")
    assert sorted(r["ResourceId"] for r in rows) == ["b-1", "i-1", "snap-1", "vol-1"]


def _recording(calls):
    def fake(region, *args):
        calls.append((region, *args[:2]))
        return fake_aws(region, *args)
    return fake


def test_completed_run_extracts_again_from_api(mod, tmp_path, monkeypatch):
    """Un nuovo run dopo uno completato NON rigenera dai file: richiama le API di Resource Explorer."""
    mod.main(_argv(tmp_path))
    calls = []
    monkeypatch.setattr(mod.elr, "aws_json", _recording(calls))
    mod.main(_argv(tmp_path))
    listed = {c[0] for c in calls if c[1:] == ("resource-explorer-2", "list-resources")}
    assert listed == {"eu-south-1", "eu-west-1", "us-east-1"}
    run = json.loads((tmp_path / ACCOUNT / "resource_explorer" / "_run.json").read_text(encoding="utf-8"))
    assert run["finished_at"] and run["regions"]["eu-south-1"] == "done"


def test_denied_region_not_retried_on_resume(mod, tmp_path, monkeypatch):
    mod.main(_argv(tmp_path))
    calls = []
    monkeypatch.setattr(mod.elr, "aws_json", _recording(calls))
    mod.main(_argv(tmp_path, "--resume"))
    assert [c for c in calls if c[1:] == ("resource-explorer-2", "list-resources")] == []
    mod.main(_argv(tmp_path, "--resume", "--retry-denied"))
    assert ("eu-west-1", "resource-explorer-2", "list-resources") in calls
