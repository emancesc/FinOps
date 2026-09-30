"""Export every resource indexed by AWS Resource Explorer into ONE xlsx per tenant.

Usage: python scripts/export_resource_explorer_xlsx.py --profile <profile>
           [--account <id>] [--tenant <name>] [--regions all] [--out-root extracted]
           [--fresh | --resume] [--retry-denied]
  --account: default = account of the profile (sts get-caller-identity).
  --tenant:  label used in the file name and summary (default: the account id).
  --regions: "all" (default) = every region with a Resource Explorer index,
             or a single region / comma-separated list.
Output: <out-root>/<account>/resource_explorer_<tenant>.xlsx with sheets
  Resources  one row per resource: ExtractionRegion (region whose index
             returned it), Region (region of the resource, "global" for
             global resources such as IAM), Service,
             ResourceType, CfnResourceType, Arn, ResourceId, Name,
             OwningAccountId, LastReportedAt, TagsLastReportedAt,
             cineca mandatory compliance, one column per cineca:* tag
             (strategy order first), one column per every other tag key,
             all tags as JSON and every non-tag property as JSON
  Summary    counts per region / resource type and cineca:* tag coverage
  Regions    extraction status per region (done / denied / error / no_index)

Resource Explorer has no aggregator index on every account, so each region
is read from its own local index (list-resources, default view).

Robustness / progress (same rules as extract_linked_resources.py):
- every run calls the Resource Explorer API: if the previous run completed,
  a new run extracts everything again from AWS. Saved pages are reused ONLY
  to resume a run that was interrupted (or explicitly with --resume).
- every page (up to 1000 resources) is appended to
  <account>/resource_explorer/<region>.jsonl as soon as it arrives, and the
  pagination token is saved in <region>.state.json: an interrupted run resumes
  from the last page; the xlsx is rebuilt from the saved pages at the end.
- permanent denials (SCP / IAM) are marked "denied" and not retried unless
  --retry-denied; other errors are retried when resuming.
- live progress: one line per page with the percentage of regions completed;
  state readable at any time in <account>/resource_explorer/_run.json.
"""
import argparse
import json
import os
import sys
import threading
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(REPO, "agent1_resource_extractor"))

import extract_linked_resources as elr  # noqa: E402  (aws_json, write_json, read_json, denied patterns)
from app.tagging_strategy import ALL_TAGS, MANDATORY_TAGS  # noqa: E402

HOME_REGION = elr.HOME_REGION
MAX_WORKERS = 6
PAGE_SIZE = 1000
MAX_CELL = 32000  # limite Excel: 32767 caratteri per cella
FIXED_COLUMNS = [
    "ExtractionRegion", "Region", "Service", "ResourceType", "CfnResourceType", "Arn", "ResourceId", "Name",
    "OwningAccountId", "LastReportedAt", "TagsLastReportedAt",
    "cineca_mandatory_compliant", "cineca_mandatory_missing",
]


def _now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


# ---------------------------------------------------------------------------
# Extraction (per region, page by page, persisted and resumable)
# ---------------------------------------------------------------------------

class Run:
    def __init__(self, path, regions):
        self._lock = threading.Lock()
        self.path = path
        self.state = elr.read_json(path, {}) or {}
        self.state.setdefault("started_at", _now())
        self.state.pop("finished_at", None)
        self.state["regions"] = {r: self.state.get("regions", {}).get(r, "pending") for r in regions}
        self.total = len(regions)
        self._save()

    @property
    def pct(self):
        final = ("done", "denied", "no_index", "error")
        finished = sum(1 for s in self.state["regions"].values() if s in final)
        return 100.0 * finished / self.total if self.total else 100.0

    def _save(self):
        self.state.update({"progress_pct": round(self.pct, 1), "updated_at": _now()})
        elr.write_json(self.path, self.state)

    def log(self, message):
        with self._lock:
            print(f"[{self.pct:5.1f}%] {message}", flush=True)

    def region(self, region, status, message=""):
        with self._lock:
            self.state["regions"][region] = status
            self._save()
            print(f"[{self.pct:5.1f}%] {region:15s} {status}{'  ' + message if message else ''}", flush=True)


def indexed_regions():
    indexes = elr.aws_json(HOME_REGION, "resource-explorer-2", "list-indexes").get("Indexes", [])
    return sorted({i["Region"] for i in indexes})


def extract_region(region, data_dir, run, fresh=False, retry_denied=False):
    jsonl = os.path.join(data_dir, f"{region}.jsonl")
    state_path = os.path.join(data_dir, f"{region}.state.json")
    state = {} if fresh else (elr.read_json(state_path, {}) or {})
    keep = ("done",) if retry_denied else ("done", "denied")
    if state.get("status") in keep:
        run.region(region, state["status"], f"{state.get('resources', 0)} resources (from disk)")
        return
    if not state.get("next_token"):
        # Nessuna pagina da cui riprendere: si riparte dall'inizio della regione
        state = {"status": "running", "pages": 0, "resources": 0, "started_at": _now()}
        open(jsonl, "w", encoding="utf-8").close()
    else:
        run.log(f"{region:15s} resume from page {state['pages'] + 1} ({state['resources']} resources on disk)")
    state["status"] = "running"
    elr.write_json(state_path, state)
    run.region(region, "running")

    while True:
        args = ["resource-explorer-2", "list-resources", "--max-results", str(PAGE_SIZE), "--no-paginate"]
        if state.get("next_token"):
            args += ["--next-token", state["next_token"]]
        try:
            page = elr.aws_json(region, *args)
        except RuntimeError as e:
            message = str(e)
            if "expired" in message.lower() or "InvalidNextToken" in message:
                state["next_token"] = None  # token scaduto: la regione ripartira' da capo
            state["status"] = "denied" if elr.is_denied_message(message) else "error"
            state["error"] = message
            elr.write_json(state_path, state)
            run.region(region, state["status"], message.splitlines()[-1][:160])
            return
        resources = page.get("Resources", [])
        with open(jsonl, "a", encoding="utf-8") as f:
            for r in resources:
                f.write(json.dumps(r, ensure_ascii=False) + "\n")
        state["pages"] += 1
        state["resources"] += len(resources)
        state["next_token"] = page.get("NextToken")
        state["view_arn"] = page.get("ViewArn")
        elr.write_json(state_path, state)
        run.log(f"{region:15s} page {state['pages']:<4} +{len(resources):<5} ({state['resources']} resources)")
        if not state["next_token"]:
            break

    state["status"] = "done"
    state["finished_at"] = _now()
    state.pop("error", None)
    elr.write_json(state_path, state)
    run.region(region, "done", f"{state['resources']} resources")


# ---------------------------------------------------------------------------
# xlsx
# ---------------------------------------------------------------------------

def _tags(resource):
    for prop in resource.get("Properties") or []:
        if prop.get("Name") == "tags":
            data = prop.get("Data") or []
            return {t.get("Key"): t.get("Value") for t in data if isinstance(t, dict)}, prop.get("LastReportedAt")
    return {}, None


def _resource_id(arn):
    tail = arn.split(":", 5)[-1] if arn.count(":") >= 5 else arn
    return tail.split("/", 1)[-1] if "/" in tail else tail


def _cell(value):
    if value is None:
        return None
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    value = str(value)
    return value if len(value) <= MAX_CELL else value[:MAX_CELL] + "...[truncated]"


def load_resources(data_dir, regions):
    rows, seen = [], set()
    for region in regions:
        path = os.path.join(data_dir, f"{region}.jsonl")
        if not os.path.exists(path):
            continue
        with open(path, encoding="utf-8") as f:
            for line in f:
                try:
                    r = json.loads(line)
                except ValueError:
                    continue  # riga troncata da un'interruzione
                if r.get("Arn") in seen:
                    continue
                seen.add(r.get("Arn"))
                r["_extraction_region"] = region
                rows.append(r)
    return rows


def build_xlsx(path, resources, region_states, account, tenant):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    tag_keys = Counter()
    parsed = []
    for r in resources:
        tags, tags_at = _tags(r)
        tag_keys.update(tags.keys())
        parsed.append((r, tags, tags_at))

    cineca_cols = [t for t in ALL_TAGS] + sorted(k for k in tag_keys if k.startswith("cineca:") and k not in ALL_TAGS)
    other_cols = sorted(k for k in tag_keys if not k.startswith("cineca:"))
    header = FIXED_COLUMNS + cineca_cols + other_cols + ["Tags (JSON)", "OtherProperties (JSON)"]

    wb = Workbook(write_only=True)
    bold, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
    cineca_fill = PatternFill("solid", fgColor="7B2C8C")

    def header_row(ws, names, special=()):
        from openpyxl.cell import WriteOnlyCell
        cells = []
        for name in names:
            c = WriteOnlyCell(ws, value=name)
            c.font, c.fill = bold, (cineca_fill if name in special else fill)
            cells.append(c)
        ws.append(cells)

    ws = wb.create_sheet("Resources")
    ws.freeze_panes = "G2"
    ws.auto_filter.ref = f"A1:{_col_letter(len(header))}{len(parsed) + 1}"
    header_row(ws, header, special=set(cineca_cols))
    coverage = Counter()
    by_region, by_type, compliant = Counter(), Counter(), 0
    for r, tags, tags_at in parsed:
        missing = [t for t in MANDATORY_TAGS if not tags.get(t)]
        compliant += not missing
        coverage.update(t for t in cineca_cols if tags.get(t))
        by_region[r.get("Region")] += 1
        by_type[r.get("ResourceType")] += 1
        others = [p for p in (r.get("Properties") or []) if p.get("Name") != "tags"]
        ws.append([
            r.get("_extraction_region"), r.get("Region"), r.get("Service"), r.get("ResourceType"), r.get("CfnResourceType"),
            r.get("Arn"), _resource_id(r.get("Arn", "")), tags.get("Name"), r.get("OwningAccountId"),
            r.get("LastReportedAt"), tags_at, "SI" if not missing else "NO", ", ".join(missing) or None,
            *[_cell(tags.get(k)) for k in cineca_cols],
            *[_cell(tags.get(k)) for k in other_cols],
            _cell(tags) if tags else None, _cell(others) if others else None,
        ])

    total = len(parsed)
    ws = wb.create_sheet("Summary")
    for row in (
        ["Tenant", tenant], ["Account", account], ["Generated at", _now()],
        ["Source", "AWS Resource Explorer (list-resources, local index per region)"],
        ["Total resources", total],
        ["cineca mandatory compliant", compliant],
        ["cineca mandatory compliant %", round(100.0 * compliant / total, 1) if total else 0.0],
        [],
    ):
        ws.append(row)
    header_row(ws, ["cineca tag", "resources with tag", "coverage %"])
    for t in cineca_cols:
        ws.append([t, coverage[t], round(100.0 * coverage[t] / total, 1) if total else 0.0])
    ws.append([])
    header_row(ws, ["Region (risorsa)", "resources"])
    for region, n in sorted(by_region.items()):
        ws.append([region, n])
    ws.append([])
    header_row(ws, ["ResourceType", "resources"])
    for rtype, n in by_type.most_common():
        ws.append([rtype, n])

    ws = wb.create_sheet("Regions")
    header_row(ws, ["Region", "status", "resources", "pages", "error"])
    for region, st in sorted(region_states.items()):
        ws.append([region, st.get("status"), st.get("resources"), st.get("pages"), _cell(st.get("error"))])

    tmp = f"{path}.tmp.xlsx"
    wb.save(tmp)
    os.replace(tmp, path)
    return total, compliant


def _col_letter(n):
    s = ""
    while n:
        n, rem = divmod(n - 1, 26)
        s = chr(65 + rem) + s
    return s


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main(argv=None):
    parser = argparse.ArgumentParser(description="Export AWS Resource Explorer resources to one xlsx per tenant")
    parser.add_argument("--profile", required=True, help="AWS CLI profile")
    parser.add_argument("--account", default=None, help="account id (default: from sts get-caller-identity)")
    parser.add_argument("--tenant", default=None, help="tenant label for file name / summary (default: account id)")
    parser.add_argument("--regions", default="all", help='"all" (default) = regions with an index, or "a,b"')
    parser.add_argument("--out-root", default=os.path.join(REPO, "extracted"), help="output root folder")
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--fresh", action="store_true", help="always extract everything again from the API")
    mode.add_argument("--resume", action="store_true",
                      help="reuse saved pages even if the previous run completed (e.g. only rebuild the xlsx)")
    parser.add_argument("--retry-denied", action="store_true", help="retry regions denied by SCP/IAM")
    args = parser.parse_args(argv)
    elr._profile = args.profile

    account = args.account or elr.aws_json(HOME_REGION, "sts", "get-caller-identity")["Account"]
    tenant = args.tenant or account
    with_index = indexed_regions()
    if args.regions.strip().lower() == "all":
        regions, no_index = with_index, []
    else:
        requested = elr.resolve_regions(args.regions)
        regions = [r for r in requested if r in with_index]
        no_index = [r for r in requested if r not in with_index]

    data_dir = os.path.join(args.out_root, account, "resource_explorer")
    os.makedirs(data_dir, exist_ok=True)
    run_path = os.path.join(data_dir, "_run.json")
    previous = elr.read_json(run_path, {}) or {}
    # Run precedente completato -> nuova estrazione completa dalle API;
    # run interrotto -> ripresa dalle pagine salvate (o --resume esplicito)
    fresh = args.fresh or (bool(previous.get("finished_at")) and not args.resume)
    if fresh and os.path.exists(run_path):
        os.remove(run_path)
    run = Run(run_path, regions + no_index)
    run.log(f"Account {account} (tenant {tenant}): {len(regions)} regions with a Resource Explorer index"
            f" -> {', '.join(regions)}")
    run.log(f"Progress file: {run_path}")
    if fresh:
        run.log("Mode: full extraction from the Resource Explorer API"
                + (f" (previous run completed at {previous['finished_at']})" if previous.get("finished_at") else ""))
    else:
        run.log("Mode: resume of an interrupted run (saved pages are reused)")
    for r in no_index:
        run.region(r, "no_index", "Resource Explorer non attivo in questa regione")

    if regions:
        with ThreadPoolExecutor(max_workers=min(MAX_WORKERS, len(regions))) as pool:
            list(pool.map(lambda r: extract_region(r, data_dir, run, fresh, args.retry_denied), regions))

    region_states = {r: elr.read_json(os.path.join(data_dir, f"{r}.state.json"), {}) or {} for r in regions}
    region_states.update({r: {"status": "no_index"} for r in no_index})
    resources = load_resources(data_dir, regions)
    xlsx = os.path.join(args.out_root, account, f"resource_explorer_{tenant}.xlsx")
    total, compliant = build_xlsx(xlsx, resources, region_states, account, tenant)
    run.state["xlsx"] = xlsx
    run.state["finished_at"] = _now()
    run._save()

    print("", flush=True)
    for region, st in sorted(region_states.items()):
        print(f"[{region}] {st.get('status')}: {st.get('resources', 0)} resources"
              + (f"  ({st['error'].splitlines()[-1][:120]})" if st.get("error") else ""), flush=True)
    print(f"\n{total} resources ({compliant} compliant with cineca mandatory tags) -> {xlsx}", flush=True)
    return xlsx


if __name__ == "__main__":
    main()
